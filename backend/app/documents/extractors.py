import asyncio,io,re,unicodedata
from dataclasses import dataclass
from docx import Document as Docx
from pypdf import PdfReader
from pptx import Presentation
from PIL import Image,ImageEnhance,ImageFilter,ImageOps
import pytesseract
from pytesseract import Output
from app.config.settings import get_settings
@dataclass
class Unit:text:str;page:int|None=None;slide:int|None=None;section:str|None=None;paragraph:int|None=None;kind:str='paragraph';confidence:float|None=None
@dataclass
class Extracted:units:list[Unit];metadata:dict
def norm(value):return re.sub(r'\n{3,}','\n\n',unicodedata.normalize('NFKC',value).replace('\x00','')).strip()
def preprocess(image):
 image=ImageOps.exif_transpose(image).convert('L');image=ImageOps.autocontrast(image);image=ImageEnhance.Contrast(image).enhance(1.4);image=image.filter(ImageFilter.MedianFilter(3))
 try:
  osd=pytesseract.image_to_osd(image,output_type=Output.DICT);rotation=int(osd.get('rotate',0));image=image.rotate(rotation,expand=True,fillcolor=255) if rotation else image
 except Exception:pass
 return image
def ocr_sync(image):
 processed=preprocess(image);data=pytesseract.image_to_data(processed,lang=get_settings().ocr_languages,config='--psm 6',output_type=Output.DICT);words=[];confidences=[]
 for text_value,confidence in zip(data.get('text',[]),data.get('conf',[])):
  if str(text_value).strip():words.append(str(text_value).strip())
  try:
   value=float(confidence)
   if value>=0:confidences.append(value/100)
  except Exception:pass
 text_value=norm(' '.join(words));confidence=sum(confidences)/len(confidences) if confidences else 0;return text_value,confidence
async def extract(data,mime):
 units=[];meta={'mime_type':mime,'ocr_pages':[],'page_quality':[]}
 if mime=='application/pdf':
  reader=PdfReader(io.BytesIO(data))
  if len(reader.pages)>get_settings().max_pdf_pages:raise ValueError('PDF has too many pages')
  meta['page_count']=len(reader.pages);fitz_doc=None
  for index,page in enumerate(reader.pages,1):
   text_value=norm(page.extract_text() or '');confidence=1.0;ocr_used=False
   if len(text_value)<40:
    try:
     import fitz
     fitz_doc=fitz_doc or fitz.open(stream=data,filetype='pdf');pix=fitz_doc[index-1].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False);text_value,confidence=await asyncio.to_thread(ocr_sync,Image.open(io.BytesIO(pix.tobytes('png'))));ocr_used=True
    except Exception:confidence=0
   if text_value:units.append(Unit(text_value,page=index,kind='page_ocr' if ocr_used else 'page',paragraph=index,confidence=confidence))
   if ocr_used:meta['ocr_pages'].append(index)
   meta['page_quality'].append({'page':index,'ocr':ocr_used,'confidence':round(confidence,4),'low_confidence':confidence<get_settings().ocr_min_confidence})
 elif mime.endswith('wordprocessingml.document'):
  document=Docx(io.BytesIO(data))
  for index,paragraph in enumerate(document.paragraphs):
   if paragraph.text.strip():kind='heading' if paragraph.style and paragraph.style.name.lower().startswith('heading') else 'paragraph';units.append(Unit(norm(paragraph.text),paragraph=index,section=paragraph.text if kind=='heading' else None,kind=kind,confidence=1))
  for index,table in enumerate(document.tables):units.append(Unit('\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows),kind='table',paragraph=index,confidence=1))
 elif mime.endswith('presentationml.presentation'):
  presentation=Presentation(io.BytesIO(data));meta['slides']=len(presentation.slides)
  for number,slide in enumerate(presentation.slides,1):
   for index,shape in enumerate(slide.shapes):
    if hasattr(shape,'text') and shape.text.strip():units.append(Unit(norm(shape.text),slide=number,section=shape.text if index==0 else None,paragraph=index,kind='slide',confidence=1))
 elif mime in {'text/plain','text/markdown'}:
  for index,value in enumerate(data.decode('utf-8',errors='replace').split('\n\n')):
   if value.strip():units.append(Unit(norm(value),paragraph=index,kind='heading' if value.lstrip().startswith('#') else 'paragraph',confidence=1))
 elif mime.startswith('image/'):
  text_value,confidence=await asyncio.to_thread(ocr_sync,Image.open(io.BytesIO(data)))
  if text_value:units=[Unit(text_value,page=1,kind='image_ocr',confidence=confidence)];meta['ocr_pages']=[1];meta['page_quality']=[{'page':1,'ocr':True,'confidence':round(confidence,4),'low_confidence':confidence<get_settings().ocr_min_confidence}]
 else:raise ValueError('Unsupported document type')
 if not units:raise ValueError('No readable text extracted')
 sample=' '.join(unit.text for unit in units)[:4000];meta['language']='hi' if sum('\u0900'<=character<='\u097f' for character in sample)>10 else 'en';confidences=[unit.confidence for unit in units if unit.confidence is not None];meta['ocr_confidence']=round(sum(confidences)/len(confidences),4) if confidences else None;meta['low_confidence']=any(item.get('low_confidence') for item in meta['page_quality']);meta['handwriting_supported']=False;return Extracted(units,meta)
def extract_text(stream,mime):
 extracted=asyncio.run(extract(stream.read(),mime));return '\n\n'.join(unit.text for unit in extracted.units),extracted.metadata
